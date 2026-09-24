#!/usr/bin/env python3
"""
Convert the uncertainty-flagging taxonomy spreadsheet into data/taxonomy.js.

    pip install openpyxl
    python3 scripts/import-taxonomy.py path/to/uncertainty_flagging_taxonomy.xlsx

Sheets read:
  Rules              code, issue type, dimension, sub-category, category, issue, fields, scope
  Temporal Keywords  keyword, example, code
  Contested words    word, note ("false positive" marks words to ignore)
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parent.parent / "data" / "taxonomy.js"


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""


def rows(ws):
    for row in ws.iter_rows(values_only=True):
        cells = [clean(c) for c in row]
        if any(cells):
            yield cells


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = {name.strip().lower(): wb[name] for name in wb.sheetnames}

    rules = []
    for cells in list(rows(sheets["rules"]))[1:]:
        code, issue_type, dimension, sub, category, issue, fields, scope = (cells + [""] * 8)[:8]
        rules.append({
            "id": re.sub(r"\s+", "", code).upper(),
            "code": code,
            "type": issue_type.lower(),
            "dimension": dimension.lower(),
            "subCategory": "" if sub in ("/", "") else sub,
            "category": category,
            "issue": issue,
            "fields": [f.strip() for f in fields.split(",") if f.strip()],
            "scope": scope,
        })

    keywords = []
    for cells in list(rows(sheets["temporal keywords"]))[1:]:
        keyword, example, code = (cells + [""] * 3)[:3]
        keywords.append({"keyword": keyword, "example": example, "rule": re.sub(r"\s+", "", code).upper()})

    contested, false_positives = [], []
    if "contested words" in sheets:
        for cells in rows(sheets["contested words"]):
            word, note = (cells + [""])[:2]
            if word.lower() in ("word", "contested word"):
                continue
            (false_positives if "false positive" in note.lower() else contested).append(word)

    data = {
        "source": re.sub(r"^[0-9a-f]{8}-", "", Path(path).name),
        "rules": rules,
        "temporalKeywords": keywords,
        "contestedWords": contested,
        "falsePositives": false_positives,
    }
    OUT.write_text(
        "/*\n * Uncertainty-flagging taxonomy, generated from "
        f"{re.sub(r'^[0-9a-f]{8}-', '', Path(path).name)}\n * by scripts/import-taxonomy.py. Do not edit by hand: update the spreadsheet\n"
        " * and re-run the script.\n */\n"
        "(function (root, taxonomy) {\n"
        "  if (typeof module === 'object' && module.exports) module.exports = taxonomy;\n"
        "  else root.MAPGAPS_TAXONOMY = taxonomy;\n"
        "})(typeof self !== 'undefined' ? self : this, "
        + json.dumps(data, indent=2, ensure_ascii=False)
        + ");\n",
        encoding="utf-8",
    )
    print(f"Wrote {len(rules)} rules, {len(keywords)} temporal keywords, "
          f"{len(contested)} contested words, {len(false_positives)} false positives to {OUT}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
